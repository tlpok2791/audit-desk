"""JN-100 분개장 편집 — 어떤 서식이든 표준 형태로 펴고 차대 일치를 검증한다."""
import io
import re
from datetime import date

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import journal as J
from core.env import in_browser, size_guard
from core.theme import rule

FONT = "Arial"
HDR = PatternFill("solid", fgColor="16202E")
SUB = PatternFill("solid", fgColor="E8E6DF")
OK_F = PatternFill("solid", fgColor="D1FAE5")
NG_F = PatternFill("solid", fgColor="FEE2E2")
_t = Side(style="thin", color="C9C6BA")
_m = Side(style="medium", color="16202E")
BOX = Border(left=_t, right=_t, top=_t, bottom=_t)


def render():
    up = st.file_uploader("분개장 — 어떤 서식이든 가능 (엑셀 · CSV)",
                          type=["xlsx", "xls", "xlsm", "csv"])
    if up is None:
        st.info("회계프로그램에서 뽑은 분개장을 그대로 올리시면 됩니다. "
                "병합셀·소계행·반복 머리글이 있어도 알아서 정리합니다.\n\n"
                "samples/ 폴더에 서로 다른 서식의 시험용 파일 두 개가 있습니다.")
        return

    lvl = size_guard(up)
    if lvl == "hard":
        st.error(
            f"파일이 {up.size/1024/1024:.0f}MB입니다. 브라우저에서 처리하기에는 너무 큽니다.\n\n"
            "이 크기는 PC에 설치해서 쓰시는 편이 훨씬 빠릅니다. "
            "홈 화면의 내려받기 안내를 참고해 주세요.")
        return
    if lvl == "warn":
        st.warning(
            f"파일이 {up.size/1024/1024:.0f}MB입니다. 브라우저에서는 몇 분 걸릴 수 있고, "
            "처리 중에는 화면이 멈춘 것처럼 보입니다. 탭을 닫지 말고 기다려 주세요.")

    # ── 읽기 ────────────────────────────────────────────────
    is_excel = up.name.lower().endswith((".xlsx", ".xls", ".xlsm"))
    try:
        if is_excel:
            xl = pd.ExcelFile(up)
            sheet = xl.sheet_names[0]
            if len(xl.sheet_names) > 1:
                sheet = st.selectbox("시트 선택", xl.sheet_names)
            probe = pd.read_excel(xl, sheet_name=sheet, header=None, nrows=20)
            guess_row = J.find_header_row(probe)
        else:
            probe, guess_row = None, 0
            for enc in ("utf-8-sig", "cp949", "utf-8"):
                try:
                    up.seek(0)
                    probe = pd.read_csv(up, encoding=enc, header=None, nrows=20)
                    enc_used = enc
                    break
                except UnicodeDecodeError:
                    continue
            if probe is None:
                st.error("인코딩을 인식하지 못했습니다. UTF-8 또는 CP949로 저장해 주세요.")
                return
            guess_row = J.find_header_row(probe)
    except Exception as e:
        st.error(f"파일을 읽지 못했습니다 — {e}")
        return

    with st.expander("원본 미리보기 (위 20행)", expanded=guess_row not in (0, -1)):
        st.dataframe(probe, use_container_width=True, height=240)

    c1, c2 = st.columns([1, 2])
    with c1:
        hdr_row = st.number_input(
            "머리글이 있는 행", min_value=1, value=(guess_row + 1) if guess_row >= 0 else 1,
            help="회사명·기간이 위에 붙어 있으면 실제 컬럼명이 있는 행 번호를 지정하세요")
    with c2:
        if guess_row >= 0:
            st.caption(f"자동으로 {guess_row + 1}행을 머리글로 판단했습니다. 다르면 왼쪽에서 고쳐주세요.")
        else:
            st.caption("머리글을 자동으로 찾지 못했습니다. 미리보기를 보고 직접 지정해 주세요.")

    try:
        if is_excel:
            raw = pd.read_excel(up, sheet_name=sheet, header=int(hdr_row) - 1)
        else:
            up.seek(0)
            raw = pd.read_csv(up, encoding=enc_used, header=int(hdr_row) - 1)
    except Exception as e:
        st.error(f"머리글 행이 맞지 않는 것 같습니다 — {e}")
        return

    raw = raw.loc[:, ~raw.columns.astype(str).str.startswith("Unnamed")]
    if raw.empty:
        st.error("데이터가 없습니다. 머리글 행 번호를 확인해 주세요.")
        return

    rule()

    # ── 컬럼 연결 ────────────────────────────────────────────
    st.markdown('<div class="section-label">컬럼 연결</div>', unsafe_allow_html=True)
    auto = J.guess_mapping(raw.columns)
    st.caption(f"{len(raw.columns)}개 컬럼 중 {len(auto)}개를 자동으로 연결했습니다. "
               "없는 항목은 '(없음)'으로 두시면 됩니다.")

    cols = ["(없음)"] + [str(c) for c in raw.columns]
    col_lookup = {str(c): c for c in raw.columns}

    amount_mode = st.radio(
        "금액 서식",
        ["차변·대변 두 컬럼", "금액 한 컬럼 (차대구분 별도)", "금액 한 컬럼 (부호로 구분)"],
        index=0 if ("차변" in auto and "대변" in auto) else (
            1 if "차대구분" in auto else 2),
        horizontal=True,
    )

    fields = ["전표일자", "전표번호", "계정코드", "계정과목", "거래처", "적요", "작성자"]
    fields += (["차변", "대변"] if amount_mode.startswith("차변")
               else (["금액", "차대구분"] if "차대구분" in amount_mode else ["금액"]))

    mapping = {}
    grid = st.columns(4)
    for i, f in enumerate(fields):
        with grid[i % 4]:
            default = str(auto[f]) if f in auto else "(없음)"
            idx = cols.index(default) if default in cols else 0
            pick = st.selectbox(f, cols, index=idx, key=f"jn_{f}")
            if pick != "(없음)":
                mapping[f] = col_lookup[pick]

    need = ["계정과목"] + (["차변", "대변"] if amount_mode.startswith("차변") else ["금액"])
    missing = [f for f in need if f not in mapping]
    if missing:
        st.error(f"필수 항목이 연결되지 않았습니다 — {', '.join(missing)}")
        return

    fill_down = st.checkbox(
        "빈 전표번호·일자를 위 값으로 채우기", value=True,
        help="병합셀로 출력된 분개장은 두 번째 줄부터 전표번호가 비어 있습니다")

    rule()

    # ── 정규화 ──────────────────────────────────────────────
    df = J.normalize(raw, mapping, fill_down=fill_down)
    if df.empty:
        st.error("정리 후 남은 데이터가 없습니다. 금액 컬럼 연결을 확인해 주세요.")
        return

    removed = len(raw) - len(df)
    m1, m2, m3 = st.columns(3)
    m1.metric("원본", f"{len(raw):,}행")
    m2.metric("정리 후", f"{len(df):,}행", f"-{removed:,}", delta_color="off")
    m3.metric("전표 수", f"{df['전표번호'].nunique():,}건")
    st.caption("소계·합계·반복 머리글·빈 금액 행을 제외한 결과입니다.")

    # ── 표지 정보 ────────────────────────────────────────────
    st.markdown('<div class="section-label">산출물 정보</div>', unsafe_allow_html=True)
    d1, d2, d3, d4 = st.columns(4)
    with d1:
        company = st.text_input("회사명", "마플")
    with d2:
        default_end = (df["전표일자"].max().date()
                       if df["전표일자"].notna().any() else date(2025, 12, 31))
        fy_end = st.date_input("결산일", value=default_end)
    with d3:
        version = st.number_input("버전", min_value=1, value=1, step=1)
    with d4:
        prepared = st.text_input("작성자", "")

    fname = J.suggest_filename(company, pd.Timestamp(fy_end), int(version))
    st.caption(f"파일명 → **{fname}**")

    rule()

    # ── 검증 ────────────────────────────────────────────────
    st.markdown('<div class="section-label">차대 검증</div>', unsafe_allow_html=True)
    res = J.verify(df, pd.Timestamp(fy_end))

    v1, v2, v3 = st.columns(3)
    v1.metric("차변 합계", f"{res['dr']:,.0f}")
    v2.metric("대변 합계", f"{res['cr']:,.0f}")
    v3.metric("차이", f"{res['diff']:,.0f}")

    if res["ok"]:
        st.success("차변과 대변이 일치하며 검증 항목에 이상이 없습니다.")
    elif abs(res["diff"]) < 1:
        st.warning("총계는 일치하지만 확인이 필요한 항목이 있습니다. 아래를 검토해 주세요.")
    else:
        st.error(f"차변과 대변이 {res['diff']:,.0f}원 일치하지 않습니다.")

    st.dataframe(res["checks"], use_container_width=True, hide_index=True)

    detail = {
        "대차 불일치 전표": res["unbalanced"],
        "차·대 동시 기재": res["both"],
        "음수 금액": res["neg"],
        "전표일자 인식 실패": res["bad_date"],
        "회계기간 범위 밖": res["out_range"],
        "계정과목 누락": res["no_acct"],
    }
    shown = {k: v for k, v in detail.items() if len(v)}
    if shown:
        st.markdown('<div class="section-label">확인 대상 상세</div>', unsafe_allow_html=True)
        for name, t in shown.items():
            with st.expander(f"{name} — {len(t):,}건"):
                st.dataframe(t, use_container_width=True, height=min(300, 40 + 35 * len(t)))

    # ── 결과 ────────────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">편집 결과</div>', unsafe_allow_html=True)
    st.dataframe(
        df.head(300).style.format({"차변": "{:,.0f}", "대변": "{:,.0f}",
                                   "전표일자": lambda d: f"{d:%Y-%m-%d}" if pd.notna(d) else ""}),
        use_container_width=True, height=380)
    if len(df) > 300:
        st.caption(f"화면에는 앞 300행만 표시했습니다. 내려받는 파일에는 {len(df):,}행 전부 들어갑니다.")

    xlsx = _export(df, res, {
        "company": company, "fy_end": pd.Timestamp(fy_end),
        "version": int(version), "prepared": prepared,
        "source": up.name, "raw_rows": len(raw), "removed": removed,
    })
    st.download_button(f"{fname} 내려받기", xlsx, file_name=fname,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ─────────────────────────────────────────────────────────────
def _export(df: pd.DataFrame, res: dict, meta: dict) -> bytes:
    wb = Workbook()

    # ── 검증 시트 ───────────────────────────────────────────
    ws = wb.active
    ws.title = "검증"
    ws["A1"] = f"{meta['company']} 분개장 편집 검증"
    ws["A1"].font = Font(FONT, size=14, bold=True)
    ws["A2"] = (f"결산일 {meta['fy_end']:%Y-%m-%d} · v{meta['version']} · "
                f"작성 {meta['prepared'] or '―'} {date.today():%Y-%m-%d}")
    ws["A2"].font = Font(FONT, size=9, italic=True, color="6B7280")
    ws["A3"] = (f"원본 {meta['source']} · {meta['raw_rows']:,}행 → "
                f"편집 후 {len(df):,}행 (소계·머리글 등 {meta['removed']:,}행 제외)")
    ws["A3"].font = Font(FONT, size=9, italic=True, color="6B7280")

    # 합계 블록
    r = 5
    ws.cell(r, 1, "구분").font = Font(FONT, bold=True, color="FFFFFF")
    ws.cell(r, 2, "금액").font = Font(FONT, bold=True, color="FFFFFF")
    for c in (1, 2):
        ws.cell(r, c).fill = HDR
        ws.cell(r, c).border = BOX
        ws.cell(r, c).alignment = Alignment("center", "center")

    last = len(df) + 1  # 편집결과 시트의 마지막 데이터 행
    for label, formula in [("차변 합계", f"=SUM('편집결과'!G2:G{last})"),
                           ("대변 합계", f"=SUM('편집결과'!H2:H{last})")]:
        r += 1
        ws.cell(r, 1, label).font = Font(FONT, size=10)
        ws.cell(r, 2, formula).number_format = "#,##0;(#,##0);-"
        for c in (1, 2):
            ws.cell(r, c).border = BOX
    r += 1
    ws.cell(r, 1, "차이 (차변 - 대변)").font = Font(FONT, bold=True)
    ws.cell(r, 2, f"=B{r-2}-B{r-1}").font = Font(FONT, bold=True)
    ws.cell(r, 2).number_format = "#,##0;(#,##0);-"
    for c in (1, 2):
        ws.cell(r, c).border = Border(top=_m, bottom=_m, left=_t, right=_t)
        ws.cell(r, c).fill = OK_F if abs(res["diff"]) < 1 else NG_F
    r += 1
    ws.cell(r, 1, "판정").font = Font(FONT, bold=True)
    ws.cell(r, 2, f'=IF(ABS(B{r-1})<1,"일치","불일치")').font = Font(FONT, bold=True)
    for c in (1, 2):
        ws.cell(r, c).border = BOX
        ws.cell(r, c).fill = OK_F if abs(res["diff"]) < 1 else NG_F

    # 검증 항목표
    r += 2
    ws.cell(r, 1, "검증 항목").font = Font(FONT, bold=True, size=11)
    r += 1
    for c, h in enumerate(["항목", "결과", "내용"], 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")
    for _, row in res["checks"].iterrows():
        r += 1
        good = row["결과"] in ("일치", "없음", "정상")
        for c, v in enumerate(row.tolist(), 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if c == 2:
                cell.fill = OK_F if good else NG_F
                cell.alignment = Alignment("center", "center")

    for col, w in zip("ABC", [24, 16, 56]):
        ws.column_dimensions[col].width = w

    # ── 편집결과 시트 ───────────────────────────────────────
    ws2 = wb.create_sheet("편집결과")
    out = df.copy()
    out["전표일자"] = out["전표일자"].dt.strftime("%Y-%m-%d").fillna("")

    for c, h in enumerate(J.STD, 1):
        cell = ws2.cell(1, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")

    for i, (_, row) in enumerate(out.iterrows(), 2):
        for c, v in enumerate(row.tolist(), 1):
            val = float(v) if c in (7, 8) else ("" if pd.isna(v) else str(v))
            cell = ws2.cell(i, c, val)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if c in (7, 8):
                cell.number_format = "#,##0;(#,##0);-"

    tot = len(out) + 2
    ws2.cell(tot, 6, "합계").font = Font(FONT, bold=True)
    for c in (7, 8):
        L = get_column_letter(c)
        ws2.cell(tot, c, f"=SUM({L}2:{L}{tot-1})")
        ws2.cell(tot, c).font = Font(FONT, bold=True)
        ws2.cell(tot, c).number_format = "#,##0"
        ws2.cell(tot, c).border = Border(top=_m, bottom=_m, left=_t, right=_t)
        ws2.cell(tot, c).fill = SUB
    ws2.cell(tot, 6).fill = SUB
    ws2.cell(tot, 6).border = Border(top=_m, bottom=_m, left=_t, right=_t)

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(J.STD))}{len(out)+1}"
    for c, w in zip("ABCDEFGHI", [12, 14, 11, 20, 18, 30, 16, 16, 10]):
        ws2.column_dimensions[c].width = w

    # ── 확인대상 시트 ───────────────────────────────────────
    flagged = {
        "대차 불일치 전표": res["unbalanced"], "차대 동시 기재": res["both"],
        "음수 금액": res["neg"], "일자 인식 실패": res["bad_date"],
        "기간 범위 밖": res["out_range"], "계정과목 누락": res["no_acct"],
    }
    flagged = {k: v for k, v in flagged.items() if len(v)}
    if flagged:
        ws3 = wb.create_sheet("확인대상")
        r = 1
        for name, t in flagged.items():
            ws3.cell(r, 1, f"{name} — {len(t):,}건").font = Font(FONT, bold=True, size=11)
            r += 1
            tt = t.copy()
            for c in tt.columns:
                if pd.api.types.is_datetime64_any_dtype(tt[c]):
                    tt[c] = tt[c].dt.strftime("%Y-%m-%d").fillna("")
            for c, h in enumerate(tt.columns, 1):
                cell = ws3.cell(r, c, str(h))
                cell.font = Font(FONT, bold=True, color="FFFFFF")
                cell.fill = HDR
                cell.border = BOX
            for _, row in tt.head(500).iterrows():
                r += 1
                for c, v in enumerate(row.tolist(), 1):
                    num = isinstance(v, (int, float)) and not isinstance(v, bool)
                    cell = ws3.cell(r, c, float(v) if num and pd.notna(v)
                                    else ("" if pd.isna(v) else str(v)))
                    cell.font = Font(FONT, size=10)
                    cell.border = BOX
                    if num:
                        cell.number_format = "#,##0;(#,##0);-"
            r += 3
        for c, w in zip("ABCDEFGHI", [14, 14, 12, 20, 18, 28, 16, 16, 10]):
            ws3.column_dimensions[c].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
