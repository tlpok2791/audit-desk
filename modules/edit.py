"""JL-100 분개장·계정별원장 편집 — 어떤 서식이든 표준 형태로 펴고 검증한다."""
import io
from datetime import date

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter

from core import journal as J
from core import ledger as L
from core.env import size_guard
from core.report import BOX, FONT, HDR, MED, NG_F, OK_F, SUB, THIN, TOTAL_BOX
from core.theme import rule

_t, _m = THIN, MED

LEDGER_OUT_COLS = ["계정과목", "날짜", "적요", "거래처코드", "거래처명", "사업자등록번호",
                   "차변금액", "대변금액", "차-대", "잔액", "계정코드", "계정명"]


def render():
    kind = st.radio(
        "편집할 자료",
        ["분개장", "계정별원장"],
        horizontal=True,
        help="분개장은 전표가 한 행씩 나열된 표, 계정별원장은 계정마다 시트가 나뉜 파일입니다.",
    )
    rule()
    if kind == "분개장":
        _render_journal()
    else:
        _render_ledger()


# ── 분개장 ═══════════════════════════════════════════════════════
def _render_journal():
    up = st.file_uploader("분개장 — 어떤 서식이든 가능 (엑셀 · CSV)",
                          type=["xlsx", "xls", "xlsm", "csv"], key="jn_up")
    if up is None:
        st.info("회계프로그램에서 뽑은 분개장을 그대로 올리시면 됩니다. "
                "병합셀·소계행·반복 머리글이 있어도 알아서 정리합니다.\n\n"
                "samples/ 폴더에 서로 다른 서식의 시험용 파일 두 개가 있습니다.")
        return

    lvl = size_guard(up)
    if lvl == "hard":
        st.error(
            f"파일이 {up.size/1024/1024:.0f}MB입니다. 브라우저에서 처리하기에는 너무 큽니다.\n\n"
            "이 크기는 PC에 설치해서 쓰시는 편이 훨씬 빠릅니다. "
            "홈 화면의 내려받기 안내를 참고해 주세요.")
        return
    if lvl == "warn":
        st.warning(
            f"파일이 {up.size/1024/1024:.0f}MB입니다. 브라우저에서는 몇 분 걸릴 수 있고, "
            "처리 중에는 화면이 멈춘 것처럼 보입니다. 탭을 닫지 말고 기다려 주세요.")

    # ── 읽기 ────────────────────────────────────────────────
    is_excel = up.name.lower().endswith((".xlsx", ".xls", ".xlsm"))
    try:
        if is_excel:
            xl = pd.ExcelFile(up)
            sheet = xl.sheet_names[0]
            if len(xl.sheet_names) > 1:
                sheet = st.selectbox("시트 선택", xl.sheet_names)
            probe = pd.read_excel(xl, sheet_name=sheet, header=None, nrows=20)
            guess_row = J.find_header_row(probe)
        else:
            probe, guess_row = None, 0
            for enc in ("utf-8-sig", "cp949", "utf-8"):
                try:
                    up.seek(0)
                    probe = pd.read_csv(up, encoding=enc, header=None, nrows=20)
                    enc_used = enc
                    break
                except UnicodeDecodeError:
                    continue
            if probe is None:
                st.error("인코딩을 인식하지 못했습니다. UTF-8 또는 CP949로 저장해 주세요.")
                return
            guess_row = J.find_header_row(probe)
    except Exception as e:
        st.error(f"파일을 읽지 못했습니다 — {e}")
        return

    with st.expander("원본 미리보기 (위 20행)", expanded=guess_row not in (0, -1)):
        st.dataframe(probe, use_container_width=True, height=240)

    c1, c2 = st.columns([1, 2])
    with c1:
        hdr_row = st.number_input(
            "머리글이 있는 행", min_value=1, value=(guess_row + 1) if guess_row >= 0 else 1,
            help="회사명·기간이 위에 붙어 있으면 실제 컬럼명이 있는 행 번호를 지정하세요")
    with c2:
        if guess_row >= 0:
            st.caption(f"자동으로 {guess_row + 1}행을 머리글로 판단했습니다. 다르면 왼쪽에서 고쳐주세요.")
        else:
            st.caption("머리글을 자동으로 찾지 못했습니다. 미리보기를 보고 직접 지정해 주세요.")

    try:
        if is_excel:
            raw = pd.read_excel(up, sheet_name=sheet, header=int(hdr_row) - 1)
        else:
            up.seek(0)
            raw = pd.read_csv(up, encoding=enc_used, header=int(hdr_row) - 1)
    except Exception as e:
        st.error(f"머리글 행이 맞지 않는 것 같습니다 — {e}")
        return

    raw = raw.loc[:, ~raw.columns.astype(str).str.startswith("Unnamed")]
    if raw.empty:
        st.error("데이터가 없습니다. 머리글 행 번호를 확인해 주세요.")
        return

    rule()

    # ── 컬럼 연결 ────────────────────────────────────────────
    st.markdown('<div class="section-label">컬럼 연결</div>', unsafe_allow_html=True)
    auto = J.guess_mapping(raw.columns)
    st.caption(f"{len(raw.columns)}개 컬럼 중 {len(auto)}개를 자동으로 연결했습니다. "
               "없는 항목은 '(없음)'으로 두시면 됩니다.")

    cols = ["(없음)"] + [str(c) for c in raw.columns]
    col_lookup = {str(c): c for c in raw.columns}

    amount_mode = st.radio(
        "금액 서식",
        ["차변·대변 두 컬럼", "금액 한 컬럼 (차대구분 별도)", "금액 한 컬럼 (부호로 구분)"],
        index=0 if ("차변" in auto and "대변" in auto) else (
            1 if "차대구분" in auto else 2),
        horizontal=True,
    )

    fields = ["전표일자", "전표번호", "계정코드", "계정과목", "거래처", "적요", "작성자"]
    fields += (["차변", "대변"] if amount_mode.startswith("차변")
               else (["금액", "차대구분"] if "차대구분" in amount_mode else ["금액"]))

    mapping = {}
    grid = st.columns(4)
    for i, f in enumerate(fields):
        with grid[i % 4]:
            default = str(auto[f]) if f in auto else "(없음)"
            idx = cols.index(default) if default in cols else 0
            pick = st.selectbox(f, cols, index=idx, key=f"jn_{f}")
            if pick != "(없음)":
                mapping[f] = col_lookup[pick]

    need = ["계정과목"] + (["차변", "대변"] if amount_mode.startswith("차변") else ["금액"])
    missing = [f for f in need if f not in mapping]
    if missing:
        st.error(f"필수 항목이 연결되지 않았습니다 — {', '.join(missing)}")
        return

    fill_down = st.checkbox(
        "빈 전표번호·일자를 위 값으로 채우기", value=True,
        help="병합셀로 출력된 분개장은 두 번째 줄부터 전표번호가 비어 있습니다")

    rule()

    # ── 정규화 ──────────────────────────────────────────────
    df = J.normalize(raw, mapping, fill_down=fill_down)
    if df.empty:
        st.error("정리 후 남은 데이터가 없습니다. 금액 컬럼 연결을 확인해 주세요.")
        return

    removed = len(raw) - len(df)
    m1, m2, m3 = st.columns(3)
    m1.metric("원본", f"{len(raw):,}행")
    m2.metric("정리 후", f"{len(df):,}행", f"-{removed:,}", delta_color="off")
    m3.metric("전표 수", f"{df['전표번호'].nunique():,}건")
    st.caption("소계·합계·반복 머리글·빈 금액 행을 제외한 결과입니다.")

    # ── 표지 정보 ────────────────────────────────────────────
    st.markdown('<div class="section-label">산출물 정보</div>', unsafe_allow_html=True)
    d1, d2, d3, d4 = st.columns(4)
    with d1:
        company = st.text_input("회사명", "마플", key="jn_company")
    with d2:
        default_end = (df["전표일자"].max().date()
                       if df["전표일자"].notna().any() else date(2025, 12, 31))
        fy_end = st.date_input("결산일", value=default_end, key="jn_fy_end")
    with d3:
        version = st.number_input("버전", min_value=1, value=1, step=1, key="jn_version")
    with d4:
        prepared = st.text_input("작성자", "", key="jn_prepared")

    fname = J.suggest_filename(company, pd.Timestamp(fy_end), int(version))
    st.caption(f"파일명 → **{fname}**")

    rule()

    # ── 검증 ────────────────────────────────────────────────
    st.markdown('<div class="section-label">차대 검증</div>', unsafe_allow_html=True)
    res = J.verify(df, pd.Timestamp(fy_end))

    v1, v2, v3 = st.columns(3)
    v1.metric("차변 합계", f"{res['dr']:,.0f}")
    v2.metric("대변 합계", f"{res['cr']:,.0f}")
    v3.metric("차이", f"{res['diff']:,.0f}")

    if res["ok"]:
        st.success("차변과 대변이 일치하며 검증 항목에 이상이 없습니다.")
    elif abs(res["diff"]) < 1:
        st.warning("총계는 일치하지만 확인이 필요한 항목이 있습니다. 아래를 검토해 주세요.")
    else:
        st.error(f"차변과 대변이 {res['diff']:,.0f}원 일치하지 않습니다.")

    st.dataframe(res["checks"], use_container_width=True, hide_index=True)

    detail = {
        "대차 불일치 전표": res["unbalanced"],
        "차·대 동시 기재": res["both"],
        "음수 금액": res["neg"],
        "전표일자 인식 실패": res["bad_date"],
        "회계기간 범위 밖": res["out_range"],
        "계정과목 누락": res["no_acct"],
    }
    shown = {k: v for k, v in detail.items() if len(v)}
    if shown:
        st.markdown('<div class="section-label">확인 대상 상세</div>', unsafe_allow_html=True)
        for name, t in shown.items():
            with st.expander(f"{name} — {len(t):,}건"):
                st.dataframe(t, use_container_width=True, height=min(300, 40 + 35 * len(t)))

    # ── 결과 ────────────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">편집 결과</div>', unsafe_allow_html=True)
    st.dataframe(
        df.head(300).style.format({"차변": "{:,.0f}", "대변": "{:,.0f}",
                                   "전표일자": lambda d: f"{d:%Y-%m-%d}" if pd.notna(d) else ""}),
        use_container_width=True, height=380)
    if len(df) > 300:
        st.caption(f"화면에는 앞 300행만 표시했습니다. 내려받는 파일에는 {len(df):,}행 전부 들어갑니다.")

    xlsx = _export_journal(df, res, {
        "company": company, "fy_end": pd.Timestamp(fy_end),
        "version": int(version), "prepared": prepared,
        "source": up.name, "raw_rows": len(raw), "removed": removed,
    })
    st.download_button(f"{fname} 내려받기", xlsx, file_name=fname,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _export_journal(df: pd.DataFrame, res: dict, meta: dict) -> bytes:
    wb = Workbook()

    # ── 검증 시트 ───────────────────────────────────────────
    ws = wb.active
    ws.title = "검증"
    ws["A1"] = f"{meta['company']} 분개장 편집 검증"
    ws["A1"].font = Font(FONT, size=14, bold=True)
    ws["A2"] = (f"결산일 {meta['fy_end']:%Y-%m-%d} · v{meta['version']} · "
                f"작성 {meta['prepared'] or '―'} {date.today():%Y-%m-%d}")
    ws["A2"].font = Font(FONT, size=9, italic=True, color="6B7280")
    ws["A3"] = (f"원본 {meta['source']} · {meta['raw_rows']:,}행 → "
                f"편집 후 {len(df):,}행 (소계·머리글 등 {meta['removed']:,}행 제외)")
    ws["A3"].font = Font(FONT, size=9, italic=True, color="6B7280")

    # 합계 블록
    r = 5
    ws.cell(r, 1, "구분").font = Font(FONT, bold=True, color="FFFFFF")
    ws.cell(r, 2, "금액").font = Font(FONT, bold=True, color="FFFFFF")
    for c in (1, 2):
        ws.cell(r, c).fill = HDR
        ws.cell(r, c).border = BOX
        ws.cell(r, c).alignment = Alignment("center", "center")

    last = len(df) + 1  # 편집결과 시트의 마지막 데이터 행
    for label, formula in [("차변 합계", f"=SUM('편집결과'!G2:G{last})"),
                           ("대변 합계", f"=SUM('편집결과'!H2:H{last})")]:
        r += 1
        ws.cell(r, 1, label).font = Font(FONT, size=10)
        ws.cell(r, 2, formula).number_format = "#,##0;(#,##0);-"
        for c in (1, 2):
            ws.cell(r, c).border = BOX
    r += 1
    ws.cell(r, 1, "차이 (차변 - 대변)").font = Font(FONT, bold=True)
    ws.cell(r, 2, f"=B{r-2}-B{r-1}").font = Font(FONT, bold=True)
    ws.cell(r, 2).number_format = "#,##0;(#,##0);-"
    for c in (1, 2):
        ws.cell(r, c).border = TOTAL_BOX
        ws.cell(r, c).fill = OK_F if abs(res["diff"]) < 1 else NG_F
    r += 1
    ws.cell(r, 1, "판정").font = Font(FONT, bold=True)
    ws.cell(r, 2, f'=IF(ABS(B{r-1})<1,"일치","불일치")').font = Font(FONT, bold=True)
    for c in (1, 2):
        ws.cell(r, c).border = BOX
        ws.cell(r, c).fill = OK_F if abs(res["diff"]) < 1 else NG_F

    # 검증 항목표
    r += 2
    ws.cell(r, 1, "검증 항목").font = Font(FONT, bold=True, size=11)
    r += 1
    for c, h in enumerate(["항목", "결과", "내용"], 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")
    for _, row in res["checks"].iterrows():
        r += 1
        good = row["결과"] in ("일치", "없음", "정상")
        for c, v in enumerate(row.tolist(), 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if c == 2:
                cell.fill = OK_F if good else NG_F
                cell.alignment = Alignment("center", "center")

    for col, w in zip("ABC", [24, 16, 56]):
        ws.column_dimensions[col].width = w

    # ── 편집결과 시트 ───────────────────────────────────────
    ws2 = wb.create_sheet("편집결과")
    out = df.copy()
    out["전표일자"] = out["전표일자"].dt.strftime("%Y-%m-%d").fillna("")

    for c, h in enumerate(J.STD, 1):
        cell = ws2.cell(1, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")

    for i, (_, row) in enumerate(out.iterrows(), 2):
        for c, v in enumerate(row.tolist(), 1):
            val = float(v) if c in (7, 8) else ("" if pd.isna(v) else str(v))
            cell = ws2.cell(i, c, val)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if c in (7, 8):
                cell.number_format = "#,##0;(#,##0);-"

    tot = len(out) + 2
    ws2.cell(tot, 6, "합계").font = Font(FONT, bold=True)
    for c in (7, 8):
        L_col = get_column_letter(c)
        ws2.cell(tot, c, f"=SUM({L_col}2:{L_col}{tot-1})")
        ws2.cell(tot, c).font = Font(FONT, bold=True)
        ws2.cell(tot, c).number_format = "#,##0"
        ws2.cell(tot, c).border = TOTAL_BOX
        ws2.cell(tot, c).fill = SUB
    ws2.cell(tot, 6).fill = SUB
    ws2.cell(tot, 6).border = TOTAL_BOX

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(J.STD))}{len(out)+1}"
    for c, w in zip("ABCDEFGHI", [12, 14, 11, 20, 18, 30, 16, 16, 10]):
        ws2.column_dimensions[c].width = w

    # ── 확인대상 시트 ───────────────────────────────────────
    flagged = {
        "대차 불일치 전표": res["unbalanced"], "차대 동시 기재": res["both"],
        "음수 금액": res["neg"], "일자 인식 실패": res["bad_date"],
        "기간 범위 밖": res["out_range"], "계정과목 누락": res["no_acct"],
    }
    flagged = {k: v for k, v in flagged.items() if len(v)}
    if flagged:
        ws3 = wb.create_sheet("확인대상")
        r = 1
        for name, t in flagged.items():
            ws3.cell(r, 1, f"{name} — {len(t):,}건").font = Font(FONT, bold=True, size=11)
            r += 1
            tt = t.copy()
            for c in tt.columns:
                if pd.api.types.is_datetime64_any_dtype(tt[c]):
                    tt[c] = tt[c].dt.strftime("%Y-%m-%d").fillna("")
            for c, h in enumerate(tt.columns, 1):
                cell = ws3.cell(r, c, str(h))
                cell.font = Font(FONT, bold=True, color="FFFFFF")
                cell.fill = HDR
                cell.border = BOX
            for _, row in tt.head(500).iterrows():
                r += 1
                for c, v in enumerate(row.tolist(), 1):
                    num = isinstance(v, (int, float)) and not isinstance(v, bool)
                    cell = ws3.cell(r, c, float(v) if num and pd.notna(v)
                                    else ("" if pd.isna(v) else str(v)))
                    cell.font = Font(FONT, size=10)
                    cell.border = BOX
                    if num:
                        cell.number_format = "#,##0;(#,##0);-"
            r += 3
        for c, w in zip("ABCDEFGHI", [14, 14, 12, 20, 18, 28, 16, 16, 10]):
            ws3.column_dimensions[c].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 계정별원장 ═══════════════════════════════════════════════════
def _render_ledger():
    up = st.file_uploader("계정별원장 — 계정마다 시트가 나뉜 파일 (.xls / .xlsx)",
                          type=["xls", "xlsx"], key="lg_up")
    if up is None:
        st.info("회계프로그램에서 뽑은 계정별원장을 그대로 올리시면 됩니다.\n\n"
                "계정과목이 시트 이름에만 있어도 열로 만들어 하나의 표로 합칩니다. "
                "[전기이월]·[월계]·[누계] 소계행은 걷어내되, "
                "적요에 쓰인 대괄호([잡코리아]유료서비스 등)는 건드리지 않습니다.")
        return

    lvl = size_guard(up)
    if lvl == "hard":
        st.error(
            f"파일이 {up.size/1024/1024:.0f}MB입니다. 브라우저에서 처리하기에는 너무 큽니다.\n\n"
            "이 크기는 PC에 설치해서 쓰시는 편이 훨씬 빠릅니다. "
            "홈 화면의 내려받기 안내를 참고해 주세요.")
        return
    if lvl == "warn":
        st.warning(
            f"파일이 {up.size/1024/1024:.0f}MB입니다. 브라우저에서는 몇 분 걸릴 수 있고, "
            "처리 중에는 화면이 멈춘 것처럼 보입니다. 탭을 닫지 말고 기다려 주세요.")

    with st.spinner("시트를 읽는 중입니다…"):
        try:
            sheets = L.read_xls_sheets(up)
        except Exception as e:
            st.error(f"파일을 읽지 못했습니다 — {e}")
            return

    if not sheets:
        st.error("계정별원장 형태의 시트를 찾지 못했습니다. "
                 "'날짜'와 '차변'이 같은 행에 있는 머리글이 필요합니다.")
        return

    st.success(f"시트 {len(sheets):,}개를 읽었습니다.")
    with st.expander("시트 인식 결과", expanded=False):
        info = pd.DataFrame([{"원본시트": s["name"], "계정코드": s["code"],
                              "계정명": s["account"], "행수": len(s["df"])}
                             for s in sheets])
        st.dataframe(info, use_container_width=True, height=280, hide_index=True)
        miss = info[info["계정코드"] == ""]
        if len(miss):
            st.warning(f"계정코드를 못 찾은 시트 {len(miss)}개 — 시트명이 그대로 계정과목이 됩니다.")

    rule()

    # ── 정리 기준 ────────────────────────────────────────────
    st.markdown('<div class="section-label">정리 기준</div>', unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    with c1:
        drop_sub = st.checkbox("소계행 제거", value=True,
                               help="[전기이월] [월계] [누계] 행을 뺍니다")
    with c2:
        drop_empty = st.checkbox("금액 0인 행 제거", value=True)
    with c3:
        sort_by = st.selectbox("정렬", ["계정코드 → 날짜", "날짜 → 계정코드", "원본 순서"])

    with st.spinner("시트를 합치는 중입니다…"):
        df, summary = L.combine(sheets, drop_subtotal=drop_sub, drop_empty=drop_empty)

    if df.empty:
        st.error("정리 후 남은 데이터가 없습니다.")
        return

    if sort_by == "계정코드 → 날짜":
        df = df.sort_values(["계정코드", "날짜"], kind="stable").reset_index(drop=True)
    elif sort_by == "날짜 → 계정코드":
        df = df.sort_values(["날짜", "계정코드"], kind="stable").reset_index(drop=True)

    raw_rows = sum(len(s["df"]) for s in sheets)
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("원본", f"{raw_rows:,}행")
    m2.metric("통합 후", f"{len(df):,}행", f"-{raw_rows-len(df):,}", delta_color="off")
    m3.metric("계정 수", f"{df['계정코드'].nunique():,}개")
    m4.metric("거래처 수", f"{df['거래처명'].replace('', pd.NA).nunique():,}개")

    # ── 산출물 정보 ──────────────────────────────────────────
    st.markdown('<div class="section-label">산출물 정보</div>', unsafe_allow_html=True)
    d1, d2, d3, d4 = st.columns(4)
    with d1:
        company = st.text_input("회사명", "마플", key="lg_company")
    with d2:
        default_end = date(2025, 12, 31)
        try:
            dd = pd.to_datetime(df["날짜"], errors="coerce", format="mixed")
            if dd.notna().any():
                default_end = dd.max().date()
        except Exception:
            pass
        period_end = st.date_input("결산일", value=default_end, key="lg_period_end")
    with d3:
        version = st.number_input("버전", min_value=1, value=1, step=1, key="lg_version")
    with d4:
        prepared = st.text_input("작성자", "", key="lg_prepared")

    fname = L.suggest_filename(company, pd.Timestamp(period_end), int(version))
    st.caption(f"파일명 → **{fname}**")

    rule()

    # ── 검증 ────────────────────────────────────────────────
    st.markdown('<div class="section-label">검증</div>', unsafe_allow_html=True)
    res = L.verify(df, summary, pd.Timestamp(period_end))

    v1, v2, v3 = st.columns(3)
    v1.metric("차변 합계", f"{res['dr']:,.0f}")
    v2.metric("대변 합계", f"{res['cr']:,.0f}")
    v3.metric("차이", f"{res['diff']:,.0f}")

    if abs(res["diff"]) < 1:
        st.success("차변과 대변이 일치합니다.")
    else:
        st.error(f"차변과 대변이 {res['diff']:,.0f}원 일치하지 않습니다. "
                 f"원장에서 빠진 계정이 있는지 확인해 주세요.")

    st.dataframe(res["checks"], use_container_width=True, hide_index=True)

    st.caption("원장 [누계]와 대조 — 원장에 찍힌 누계 금액과 편집 결과가 맞는지 계정별로 "
               "대조합니다. 소계행을 제대로 걷어냈는지 확인하는 가장 확실한 방법입니다.")

    detail = {
        "[누계] 불일치 계정": res["bad_tie"],
        "차·대 동시 기재": res["both"],
        "음수 금액": res["neg"],
        "날짜 인식 실패": res["bad_date"],
        "회계기간 범위 밖": res["out_range"],
    }
    shown = {k: v for k, v in detail.items() if len(v)}
    if shown:
        st.markdown('<div class="section-label">확인 대상</div>', unsafe_allow_html=True)
        for name, t in shown.items():
            with st.expander(f"{name} — {len(t):,}건"):
                st.dataframe(t.head(500), use_container_width=True,
                             height=min(300, 40 + 35 * len(t)))

    # ── 계정별 요약 ──────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">계정별 집계</div>', unsafe_allow_html=True)
    sm_view = summary[["계정코드", "계정명", "행수", "차변합계", "대변합계"]].copy()
    sm_view["차-대"] = sm_view["차변합계"] - sm_view["대변합계"]
    st.dataframe(
        sm_view.style.format({"차변합계": "{:,.0f}", "대변합계": "{:,.0f}", "차-대": "{:,.0f}"}),
        use_container_width=True, height=320, hide_index=True)

    # ── 결과 ────────────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">통합 원장</div>', unsafe_allow_html=True)
    cols = [c for c in LEDGER_OUT_COLS if c in df.columns]
    st.dataframe(
        df[cols].head(300).style.format(
            {"차변금액": "{:,.0f}", "대변금액": "{:,.0f}",
             "차-대": "{:,.0f}", "잔액": "{:,.0f}"}),
        use_container_width=True, height=380)
    if len(df) > 300:
        st.caption(f"화면에는 앞 300행만 표시했습니다. "
                   f"내려받는 파일에는 {len(df):,}행 전부 들어갑니다.")

    if st.button("엑셀 만들기", type="primary"):
        with st.spinner(f"{len(df):,}행을 쓰는 중입니다…"):
            xlsx = _export_ledger(df[cols], summary, res, {
                "company": company, "period_end": pd.Timestamp(period_end),
                "version": int(version), "prepared": prepared,
                "source": up.name, "raw_rows": raw_rows,
                "sheet_count": len(sheets),
            })
        st.download_button(f"{fname} 내려받기", xlsx, file_name=fname,
                           mime="application/vnd.openxmlformats-officedocument."
                                "spreadsheetml.sheet")


def _export_ledger(df: pd.DataFrame, summary: pd.DataFrame, res: dict, meta: dict) -> bytes:
    wb = Workbook()

    # ── 검증 ────────────────────────────────────────────────
    ws = wb.active
    ws.title = "검증"
    ws["A1"] = f"{meta['company']} 계정별원장 통합 검증"
    ws["A1"].font = Font(FONT, size=14, bold=True)
    ws["A2"] = (f"결산일 {meta['period_end']:%Y-%m-%d} · v{meta['version']} · "
                f"작성 {meta['prepared'] or '―'} {date.today():%Y-%m-%d}")
    ws["A2"].font = Font(FONT, size=9, italic=True, color="6B7280")
    ws["A3"] = (f"원본 {meta['source']} · 시트 {meta['sheet_count']:,}개 · "
                f"{meta['raw_rows']:,}행 → 통합 {len(df):,}행")
    ws["A3"].font = Font(FONT, size=9, italic=True, color="6B7280")

    last = len(df) + 2  # 통합원장 시트의 마지막 데이터 행 (1행 합계, 2행 머리글)
    r = 5
    for c, h in enumerate(["구분", "금액"], 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")
    for label, f in [("차변 합계", f"=SUM('통합원장'!G3:G{last})"),
                     ("대변 합계", f"=SUM('통합원장'!H3:H{last})")]:
        r += 1
        ws.cell(r, 1, label).font = Font(FONT, size=10)
        ws.cell(r, 2, f).number_format = "#,##0;(#,##0);-"
        for c in (1, 2):
            ws.cell(r, c).border = BOX
    r += 1
    ws.cell(r, 1, "차이 (차변 - 대변)").font = Font(FONT, bold=True)
    ws.cell(r, 2, f"=B{r-2}-B{r-1}").font = Font(FONT, bold=True)
    ws.cell(r, 2).number_format = "#,##0;(#,##0);-"
    good = abs(res["diff"]) < 1
    for c in (1, 2):
        ws.cell(r, c).border = TOTAL_BOX
        ws.cell(r, c).fill = OK_F if good else NG_F
    r += 1
    ws.cell(r, 1, "판정").font = Font(FONT, bold=True)
    ws.cell(r, 2, f'=IF(ABS(B{r-1})<1,"일치","불일치")').font = Font(FONT, bold=True)
    for c in (1, 2):
        ws.cell(r, c).border = BOX
        ws.cell(r, c).fill = OK_F if good else NG_F

    r += 2
    ws.cell(r, 1, "검증 항목").font = Font(FONT, bold=True, size=11)
    r += 1
    for c, h in enumerate(["항목", "결과", "내용"], 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")
    for _, row in res["checks"].iterrows():
        r += 1
        g = row["결과"] in ("일치", "없음", "정상", "참고")
        for c, v in enumerate(row.tolist(), 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if c == 2:
                cell.fill = OK_F if g else NG_F
                cell.alignment = Alignment("center", "center")
    for col, w in zip("ABC", [24, 16, 60]):
        ws.column_dimensions[col].width = w

    # ── 통합원장 — 편집본 서식 (1행 합계, 2행 머리글) ──────────
    ws2 = wb.create_sheet("통합원장")
    cols = list(df.columns)
    i_dr = cols.index("차변금액") + 1
    i_cr = cols.index("대변금액") + 1
    Ldr, Lcr = get_column_letter(i_dr), get_column_letter(i_cr)

    # 1행: 전체 범위 합계 (SUBTOTAL 아님 — 필터에 영향받지 않는다)
    ws2.cell(1, i_dr, f"=SUM({Ldr}3:{Ldr}{last})")
    ws2.cell(1, i_cr, f"=SUM({Lcr}3:{Lcr}{last})")
    for c in (i_dr, i_cr):
        cell = ws2.cell(1, c)
        cell.font = Font(FONT, bold=True, size=11)
        cell.number_format = "#,##0"
        cell.fill = SUB
        cell.border = TOTAL_BOX

    for c, h in enumerate(cols, 1):
        cell = ws2.cell(2, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")

    num_idx = {cols.index(c) + 1 for c in ("차변금액", "대변금액", "차-대", "잔액")
               if c in cols}
    body = Font(FONT, size=10)
    for i, rec in enumerate(df.itertuples(index=False), 3):
        for c, v in enumerate(rec, 1):
            cell = ws2.cell(i, c, v)
            cell.font = body
            if c in num_idx:
                cell.number_format = "#,##0;(#,##0);-"

    ws2.freeze_panes = "A3"
    ws2.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{last}"
    widths = {"계정과목": 22, "날짜": 12, "적요": 34, "거래처코드": 11, "거래처명": 26,
              "사업자등록번호": 15, "차변금액": 16, "대변금액": 16, "차-대": 16,
              "잔액": 18, "계정코드": 10, "계정명": 18}
    for c, h in enumerate(cols, 1):
        ws2.column_dimensions[get_column_letter(c)].width = widths.get(h, 14)

    # ── 계정별 집계 ──────────────────────────────────────────
    ws3 = wb.create_sheet("계정별집계")
    sm = summary[["계정코드", "계정명", "행수", "차변합계", "대변합계",
                  "전기이월_차변", "전기이월_대변", "원장누계_차변", "원장누계_대변"]].copy()
    for c, h in enumerate(list(sm.columns) + ["차-대", "누계대조"], 1):
        cell = ws3.cell(1, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")
    for i, rec in enumerate(sm.itertuples(index=False), 2):
        for c, v in enumerate(rec, 1):
            cell = ws3.cell(i, c, None if pd.isna(v) else v)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if c >= 4:
                cell.number_format = "#,##0;(#,##0);-"
        ws3.cell(i, 10, f"=D{i}-E{i}").number_format = "#,##0;(#,##0);-"
        ws3.cell(i, 11, f'=IF(H{i}="","―",IF(ABS(D{i}+F{i}-H{i})<1,"일치","불일치"))')
        for c in (10, 11):
            ws3.cell(i, c).font = Font(FONT, size=10)
            ws3.cell(i, c).border = BOX

    t = len(sm) + 2
    ws3.cell(t, 2, "합계").font = Font(FONT, bold=True)
    for c in (4, 5):
        Lc = get_column_letter(c)
        ws3.cell(t, c, f"=SUM({Lc}2:{Lc}{t-1})")
        ws3.cell(t, c).font = Font(FONT, bold=True)
        ws3.cell(t, c).number_format = "#,##0"
        ws3.cell(t, c).fill = SUB
        ws3.cell(t, c).border = TOTAL_BOX
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:K{t-1}"
    for c, w in zip("ABCDEFGHIJK", [10, 20, 9, 17, 17, 15, 15, 17, 17, 17, 11]):
        ws3.column_dimensions[c].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
