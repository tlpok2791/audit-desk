"""IT-100 종합소득세 계산 — 사업소득자용 참고용 세액 계산기.

수입금액·필요경비·공제 항목을 홈택스 엑셀·영수증 PDF·직접입력으로 채우면
core/income_tax.py 의 계산 엔진이 신고서 흐름(사업소득금액 → 과세표준 →
산출세액 → 결정세액 → 차감납부세액)을 그대로 보여준다.

참고용 추정치다. 가산세·특별세액공제(보험료·의료비·교육비 등)는 범위 밖이며
실제 신고 전 세무사 확인이 필요하다 — 화면에도 계속 띄운다.
"""
import io

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core import income_tax as IT
from core import loader, receipt
from core.env import in_browser
from core.theme import rule

FONT = "Arial"
HDR = PatternFill("solid", fgColor="16202E")
SUB = PatternFill("solid", fgColor="E8E6DF")
_t = Side(style="thin", color="C9C6BA")
BOX = Border(left=_t, right=_t, top=_t, bottom=_t)

LABELS = {
    "수입금액": "수입금액", "필요경비": "필요경비", "사업소득금액": "사업소득금액",
    "종합소득금액": "종합소득금액", "인적공제": "인적공제", "국민연금등_공제": "국민연금 등 공제",
    "종합소득공제": "종합소득공제 합계", "과세표준": "과세표준", "산출세액": "산출세액",
    "표준세액공제": "표준세액공제", "자녀세액공제": "자녀세액공제", "세액공제_합계": "세액공제 합계",
    "결정세액": "결정세액", "기납부세액": "기납부세액(원천징수)", "차감납부세액": "차감납부(환급)세액",
    "지방소득세_추정": "지방소득세(추정, 별도 신고)",
}


def _init_state():
    for key, default in [
        ("it_수입금액", 0), ("it_필요경비", 0), ("it_국민연금", 0), ("it_기납부세액", 0),
    ]:
        st.session_state.setdefault(key, default)


def _income_section() -> float:
    st.markdown('<div class="section-label">1 · 수입금액</div>', unsafe_allow_html=True)
    st.caption("홈택스에서 내려받은 매출·수입 자료 엑셀을 올리거나, 합계를 직접 입력하세요.")
    up = st.file_uploader("수입금액 자료 (선택)", type=["csv", "xlsx", "xls", "xlsm"], key="it_income_file")
    if up is not None:
        try:
            raw = loader.read_any(up)
        except Exception as e:
            st.error(f"파일을 읽지 못했습니다 — {e}")
            raw = None
        if raw is not None:
            st.caption(f"{len(raw):,}행 × {len(raw.columns)}열")
            num_cols = [c for c in raw.columns if pd.to_numeric(raw[c], errors="coerce").notna().any()]
            if num_cols:
                col = st.selectbox("합산할 금액 열", num_cols, key="it_income_col")
                total = float(pd.to_numeric(raw[col], errors="coerce").fillna(0).sum())
                st.session_state["it_수입금액"] = total
                st.info(f"'{col}' 열 합계 {total:,.0f}원을 수입금액으로 채웠습니다. 아래에서 고칠 수 있습니다.")
    return st.number_input("수입금액 (원)", min_value=0, step=100000, format="%d", key="it_수입금액")


def _expense_section() -> float:
    st.markdown('<div class="section-label">2 · 필요경비</div>', unsafe_allow_html=True)
    method = st.radio(
        "방식", ["영수증 PDF 업로드", "직접 입력 장부", "경비율로 추정"],
        horizontal=True, key="it_expense_method",
    )

    if method == "영수증 PDF 업로드":
        if in_browser():
            st.warning("PDF 파싱은 PC 설치 실행에서만 됩니다 (`streamlit run app.py`). "
                       "브라우저에서는 '직접 입력 장부'를 써주세요.")
            return st.session_state.get("it_영수증합계", 0.0)
        if not receipt.available():
            st.error("pdfplumber가 설치되어 있지 않습니다 — `pip install -r requirements.txt` 후 다시 실행하세요.")
            return 0.0
        files = st.file_uploader("경비 영수증·세금계산서 PDF (여러 개)", type=["pdf"],
                                 accept_multiple_files=True, key="it_receipt_files")
        rows = []
        for f in files or []:
            try:
                text = receipt.extract_text(f)
            except Exception as e:
                rows.append({"파일": f.name, "근거줄": f"읽기 실패 — {e}", "금액": 0})
                continue
            guess = receipt.guess_amount(text)
            rows.append({"파일": f.name, "근거줄": guess["근거줄"] if guess else "합계금액 패턴을 찾지 못함",
                        "금액": guess["금액"] if guess else 0})
        if not rows:
            st.info("PDF를 올리면 '합계금액' 류 문구 주변 숫자를 후보로 뽑아 보여줍니다. "
                    "서식마다 달라 틀릴 수 있으니 아래 표에서 직접 확인·수정하세요.")
            return 0.0
        edited = st.data_editor(pd.DataFrame(rows), use_container_width=True, hide_index=True, key="it_receipt_edit")
        total = float(pd.to_numeric(edited["금액"], errors="coerce").fillna(0).sum())
        st.session_state["it_영수증합계"] = total
        st.metric("영수증 합계", f"{total:,.0f}원")
        return total

    if method == "직접 입력 장부":
        st.caption("경비 항목을 행마다 입력하세요.")
        default = pd.DataFrame([{"항목": "", "금액": 0}])
        edited = st.data_editor(
            st.session_state.get("it_장부", default),
            num_rows="dynamic", use_container_width=True, hide_index=True, key="it_ledger_edit",
        )
        st.session_state["it_장부"] = edited
        total = float(pd.to_numeric(edited["금액"], errors="coerce").fillna(0).sum())
        st.metric("장부 합계", f"{total:,.0f}원")
        return total

    # 경비율로 추정
    st.caption("업종별 기준경비율·단순경비율은 홈택스에서 확인해 직접 입력하세요.")
    수입금액 = st.session_state.get("it_수입금액", 0)
    rate = st.number_input("경비율 (%)", min_value=0.0, max_value=100.0, value=60.0, step=0.5, key="it_rate")
    total = float(수입금액) * rate / 100
    st.metric("추정 필요경비", f"{total:,.0f}원")
    return total


def _deduction_section() -> tuple[dict, float]:
    st.markdown('<div class="section-label">3 · 인적공제 · 국민연금 등</div>', unsafe_allow_html=True)
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        기본 = st.number_input("기본공제 인원(본인 포함)", min_value=1, value=1, step=1, key="it_기본인원")
    with c2:
        경로 = st.number_input("70세 이상 인원", min_value=0, value=0, step=1, key="it_경로인원")
    with c3:
        장애 = st.number_input("장애인 인원", min_value=0, value=0, step=1, key="it_장애인원")
    with c4:
        자녀 = st.number_input("8세 이상 자녀 수(세액공제용)", min_value=0, value=0, step=1, key="it_자녀수")

    c5, c6 = st.columns(2)
    with c5:
        모드 = st.radio("부녀자 / 한부모 공제", ["해당 없음", "부녀자공제", "한부모공제"],
                       horizontal=True, key="it_부녀자모드")
    with c6:
        국민연금 = st.number_input("국민연금 등 공적연금 본인부담분 (원)", min_value=0, step=10000,
                                 format="%d", key="it_국민연금")

    inputs = {
        "기본공제_인원": 기본, "경로우대_인원": 경로, "장애인_인원": 장애,
        "부녀자_해당": 모드 == "부녀자공제", "한부모_해당": 모드 == "한부모공제",
        "자녀세액공제_인원": 자녀,
    }
    return inputs, float(국민연금)


def _withholding_section() -> float:
    st.markdown('<div class="section-label">4 · 기납부세액 (원천징수)</div>', unsafe_allow_html=True)
    if not in_browser() and receipt.available():
        f = st.file_uploader("원천징수영수증 PDF (선택)", type=["pdf"], key="it_wh_file")
        if f is not None:
            try:
                text = receipt.extract_text(f)
                guess = receipt.guess_withholding_tax(text)
            except Exception as e:
                st.error(f"읽지 못했습니다 — {e}")
                guess = None
            applied_for = st.session_state.get("it_wh_applied_for")
            if guess and applied_for != f.name:
                st.info(f"'{guess['근거줄']}' 줄에서 {guess['금액']:,.0f}원을 찾아 아래 값에 채웠습니다. "
                        "맞는지 확인하고 다르면 직접 고쳐주세요.")
                st.session_state["it_기납부세액"] = int(guess["금액"])
                st.session_state["it_wh_applied_for"] = f.name
    return st.number_input("기납부세액 합계 (원천징수영수증 소득세 합)", min_value=0, step=1000,
                           format="%d", key="it_기납부세액")


def _to_excel(result: dict, meta: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "종합소득세 계산"
    ws["A1"] = f"{meta['company']} — {meta['year']}년 귀속 종합소득세 계산(참고용)"
    ws["A1"].font = Font(FONT, size=13, bold=True)
    ws["A2"] = "실제 신고 전 세무사 확인이 필요한 추정치입니다."
    ws["A2"].font = Font(FONT, size=9, italic=True, color="B3261E")

    row = 4
    for key, label in LABELS.items():
        ws.cell(row, 1, label).font = Font(FONT, bold=key in ("과세표준", "결정세액", "차감납부세액"))
        cell = ws.cell(row, 2, result[key])
        cell.number_format = "#,##0"
        for c in (1, 2):
            ws.cell(row, c).border = BOX
        if key in ("과세표준", "결정세액", "차감납부세액"):
            ws.cell(row, 1).fill = SUB
            ws.cell(row, 2).fill = SUB
        row += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render():
    _init_state()
    st.warning("**참고용 추정치입니다.** 가산세, 보험료·의료비·교육비 등 특별세액공제는 계산에 "
              "포함하지 않습니다. 실제 신고 전 세무사 확인을 받으세요.", icon="⚠️")

    c1, c2 = st.columns(2)
    with c1:
        company = st.text_input("상호 / 성명", "")
    with c2:
        year = st.number_input("귀속연도", min_value=2020, max_value=2100, value=2025, step=1)

    rule()
    수입금액 = _income_section()
    rule()
    필요경비 = _expense_section()
    rule()
    ded_inputs, 국민연금 = _deduction_section()
    rule()
    기납부세액 = _withholding_section()
    rule()

    inputs = {
        "수입금액": 수입금액, "필요경비": 필요경비, "국민연금등_공제": 국민연금,
        "기납부세액": 기납부세액, **ded_inputs,
    }
    result = IT.calc(inputs)

    st.markdown('<div class="section-label">계산 결과</div>', unsafe_allow_html=True)
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("과세표준", f"{result['과세표준']:,.0f}원")
    m2.metric("산출세액", f"{result['산출세액']:,.0f}원")
    m3.metric("결정세액", f"{result['결정세액']:,.0f}원")
    label = "환급세액" if result["환급여부"] else "납부세액"
    m4.metric(label, f"{abs(result['차감납부세액']):,.0f}원")

    detail = pd.DataFrame(
        [{"항목": LABELS[k], "금액": v} for k, v in result.items() if k != "환급여부"]
    )
    st.dataframe(detail, use_container_width=True, hide_index=True,
                column_config={"금액": st.column_config.NumberColumn(format="%d원")})

    xlsx = _to_excel(result, {"company": company or "미입력", "year": int(year)})
    st.download_button("계산 결과 엑셀 내려받기", xlsx,
                       file_name=f"종합소득세_계산_{company or '미입력'}_{year}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
