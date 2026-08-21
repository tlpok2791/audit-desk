"""FS-200 재무제표 분석 — 증감분석 · 비율분석 · 코멘트 초안."""
import io

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import accounts, loader
from core.theme import rule

FIELDS = ["계정과목", "당기금액", "전기금액"]
FONT = "Arial"
HDR = PatternFill("solid", fgColor="16202E")
SUB = PatternFill("solid", fgColor="E8E6DF")
FLAG = PatternFill("solid", fgColor="FCE8E6")
_t = Side(style="thin", color="C9C6BA")
BOX = Border(left=_t, right=_t, top=_t, bottom=_t)


def render():
    up = st.file_uploader("시산표 — 계정과목 / 당기금액 / 전기금액",
                          type=["csv", "xlsx", "xls", "xlsm"])
    if up is None:
        st.info("2개년 시산표를 올려주세요. samples/trial_balance.csv로 시험해 보셔도 됩니다.")
        return

    try:
        raw = loader.read_any(up)
    except Exception as e:
        st.error(f"파일을 읽지 못했습니다 — {e}")
        return

    st.caption(f"{len(raw):,}행 × {len(raw.columns)}열")
    mapping = loader.map_columns(raw, FIELDS)
    if mapping is None:
        return

    df = loader.standardize_numeric(loader.rename(raw, mapping), ["당기금액", "전기금액"])
    df = df.dropna(subset=["계정과목"]).fillna({"당기금액": 0, "전기금액": 0})
    df = accounts.classify(df)

    # ── 분류 확인 ────────────────────────────────────────────
    unknown = df[df["세부구분"] == "미분류"]
    with st.expander(f"계정 분류 확인 — 자동 분류 {len(df)-len(unknown)}건 / 미분류 {len(unknown)}건",
                     expanded=len(unknown) > 0):
        st.caption("계정명으로 추정한 결과입니다. 틀린 항목은 세부구분을 직접 고쳐주세요.")
        edited = st.data_editor(
            df[["계정과목", "세부구분", "당기금액", "전기금액"]],
            column_config={"세부구분": st.column_config.SelectboxColumn(options=accounts.ORDER)},
            use_container_width=True, hide_index=True, height=320, key="cls",
        )
        df["세부구분"] = edited["세부구분"].values
        df["구분"] = np.where(df["세부구분"].isin(["유동자산", "비유동자산"]), "자산",
                    np.where(df["세부구분"].isin(["유동부채", "비유동부채"]), "부채",
                    np.where(df["세부구분"] == "자본", "자본",
                    np.where(df["세부구분"] == "미분류", "미분류", "손익"))))

    rule()

    # ── 임계치 ──────────────────────────────────────────────
    st.markdown('<div class="section-label">유의성 기준</div>', unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    with c1:
        amt_th = st.number_input("금액 임계치 (원)", value=100_000_000, step=10_000_000, format="%d",
                                 help="이 금액을 넘는 증감만 유의적으로 봅니다")
    with c2:
        pct_th = st.slider("증감률 임계치 (%)", 5, 100, 20)
    with c3:
        both = st.radio("적용 방식", ["금액 AND 증감률", "금액 OR 증감률"], index=0,
                        help="AND는 둘 다 넘을 때만, OR은 하나만 넘어도 적출")

    df["증감액"] = df["당기금액"] - df["전기금액"]
    df["증감률"] = np.where(df["전기금액"] != 0, df["증감액"] / df["전기금액"].abs(), np.nan)

    over_a = df["증감액"].abs() >= amt_th
    over_p = df["증감률"].abs() >= pct_th / 100
    df["유의"] = (over_a & over_p) if both.endswith("증감률") and "AND" in both else (over_a | over_p)

    rule()

    # ── 재무제표 요약 ────────────────────────────────────────
    cur, pri = "당기금액", "전기금액"
    S = lambda s, c: accounts.subtotal(df, s, c)
    tot_asset_c = S("유동자산", cur) + S("비유동자산", cur)
    tot_asset_p = S("유동자산", pri) + S("비유동자산", pri)
    tot_liab_c = S("유동부채", cur) + S("비유동부채", cur)
    tot_liab_p = S("유동부채", pri) + S("비유동부채", pri)
    equity_c, equity_p = S("자본", cur), S("자본", pri)
    rev_c, rev_p = S("매출", cur), S("매출", pri)
    cogs_c, cogs_p = S("매출원가", cur), S("매출원가", pri)
    sga_c, sga_p = S("판매비와관리비", cur), S("판매비와관리비", pri)
    op_c, op_p = rev_c - cogs_c - sga_c, rev_p - cogs_p - sga_p

    for label, a, l, e in [("당기", tot_asset_c, tot_liab_c, equity_c),
                           ("전기", tot_asset_p, tot_liab_p, equity_p)]:
        gap = a - l - e
        if abs(gap) > 1:
            st.warning(f"{label} 대차 불일치 {gap:,.0f}원 — 자산 {a:,.0f} vs 부채+자본 {l+e:,.0f}. "
                       f"계정 분류나 누락 계정을 먼저 확인해 주세요.")

    st.markdown('<div class="section-label">요약</div>', unsafe_allow_html=True)
    m = st.columns(4)
    for col, (label, c, p) in zip(m, [
        ("총자산", tot_asset_c, tot_asset_p), ("총부채", tot_liab_c, tot_liab_p),
        ("매출", rev_c, rev_p), ("영업이익", op_c, op_p),
    ]):
        with col:
            d = (c - p) / abs(p) * 100 if p else 0
            st.metric(label, f"{c/1e8:,.1f}억", f"{d:+.1f}%")

    # ── 증감분석 ────────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">증감분석</div>', unsafe_allow_html=True)
    flagged = df[df["유의"]].sort_values("증감액", key=abs, ascending=False)
    st.caption(f"유의적 증감 {len(flagged)}개 계정 / 전체 {len(df)}개")

    show = flagged[["세부구분", "계정과목", "전기금액", "당기금액", "증감액", "증감률"]].copy()
    st.dataframe(
        show.style.format({"전기금액": "{:,.0f}", "당기금액": "{:,.0f}",
                           "증감액": "{:,.0f}", "증감률": "{:+.1%}"}),
        use_container_width=True, height=min(420, 40 + 35 * len(show)))

    # ── 비율분석 ────────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">비율분석</div>', unsafe_allow_html=True)
    ratios = _ratios(df, S, tot_asset_c, tot_asset_p, tot_liab_c, tot_liab_p,
                     equity_c, equity_p, rev_c, rev_p, cogs_c, cogs_p, op_c, op_p)
    st.dataframe(ratios, use_container_width=True, hide_index=True)

    # ── 코멘트 초안 ──────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">코멘트 초안</div>', unsafe_allow_html=True)
    st.caption("규칙 기반 초안입니다. 그대로 쓰지 마시고 근거를 확인한 뒤 다듬어 주세요.")
    comments = _comments(df, flagged, rev_c, rev_p, cogs_c, cogs_p, sga_c, sga_p)
    for c in comments:
        st.markdown(f"- {c}")
    st.text_area("편집", "\n".join(f"- {c}" for c in comments), height=200, key="cmt")

    # ── 내려받기 ────────────────────────────────────────────
    rule()
    xlsx = _export(df, flagged, ratios, comments, amt_th, pct_th)
    st.download_button("분석조서 내려받기 (xlsx)", xlsx, file_name="FS_증감분석.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ─────────────────────────────────────────────────────────────
def _pct(n, d):
    return n / d if d else np.nan


def _ratios(df, S, ta_c, ta_p, tl_c, tl_p, eq_c, eq_p, rev_c, rev_p, cogs_c, cogs_p, op_c, op_p):
    ca_c, ca_p = S("유동자산", "당기금액"), S("유동자산", "전기금액")
    cl_c, cl_p = S("유동부채", "당기금액"), S("유동부채", "전기금액")
    inv_c = accounts.find(df, "재고", "당기금액")
    inv_p = accounts.find(df, "재고", "전기금액")
    ar_c = accounts.find(df, "매출채권", "당기금액")
    ar_p = accounts.find(df, "매출채권", "전기금액")
    int_c = accounts.find(df, "이자비용", "당기금액")
    int_p = accounts.find(df, "이자비용", "전기금액")

    items = [
        ("유동비율", _pct(ca_c, cl_c), _pct(ca_p, cl_p), "%", "유동자산 ÷ 유동부채"),
        ("당좌비율", _pct(ca_c - inv_c, cl_c), _pct(ca_p - inv_p, cl_p), "%", "(유동자산-재고) ÷ 유동부채"),
        ("부채비율", _pct(tl_c, eq_c), _pct(tl_p, eq_p), "%", "총부채 ÷ 자본"),
        ("자기자본비율", _pct(eq_c, ta_c), _pct(eq_p, ta_p), "%", "자본 ÷ 총자산"),
        ("매출총이익률", _pct(rev_c - cogs_c, rev_c), _pct(rev_p - cogs_p, rev_p), "%", "매출총이익 ÷ 매출"),
        ("영업이익률", _pct(op_c, rev_c), _pct(op_p, rev_p), "%", "영업이익 ÷ 매출"),
        ("총자산회전율", _pct(rev_c, ta_c), _pct(rev_p, ta_p), "회", "매출 ÷ 총자산"),
        ("매출채권회전율", _pct(rev_c, ar_c), _pct(rev_p, ar_p), "회", "매출 ÷ 매출채권"),
        ("재고자산회전율", _pct(cogs_c, inv_c), _pct(cogs_p, inv_p), "회", "매출원가 ÷ 재고자산"),
        ("이자보상배율", _pct(op_c, int_c), _pct(op_p, int_p), "배", "영업이익 ÷ 이자비용"),
    ]
    rows = []
    for name, c, p, unit, basis in items:
        fmt = (lambda v: "—" if pd.isna(v) else (f"{v*100:,.1f}%" if unit == "%" else f"{v:,.2f}{unit}"))
        chg = "—"
        if not (pd.isna(c) or pd.isna(p)):
            d = (c - p) * 100 if unit == "%" else c - p
            chg = f"{d:+,.1f}%p" if unit == "%" else f"{d:+,.2f}{unit}"
        rows.append({"지표": name, "당기": fmt(c), "전기": fmt(p), "증감": chg, "산식": basis})
    return pd.DataFrame(rows)


def _comments(df, flagged, rev_c, rev_p, cogs_c, cogs_p, sga_c, sga_p):
    out = []
    g = lambda c, p: (c - p) / abs(p) if p else np.nan
    F = lambda k, col: accounts.find(df, k, col)

    rev_g, cogs_g = g(rev_c, rev_p), g(cogs_c, cogs_p)

    # 매출채권 대 매출
    ar_c, ar_p = F("매출채권", "당기금액"), F("매출채권", "전기금액")
    if ar_p and not pd.isna(rev_g):
        ar_g = g(ar_c, ar_p)
        if ar_g - rev_g > 0.10:
            out.append(
                f"**매출채권**이 전기 대비 {ar_g:+.1%}({ar_c-ar_p:,.0f}원) 변동하여 "
                f"매출 증가율 {rev_g:+.1%}를 {(ar_g-rev_g)*100:.0f}%p 상회함. "
                f"회수기간 장기화 또는 기말 매출인식 시점 검토가 필요함.")
        # 대손충당금 설정률
        al_c, al_p = abs(F("대손충당", "당기금액")), abs(F("대손충당", "전기금액"))
        if ar_c and al_c:
            r_c, r_p = al_c / ar_c, (al_p / ar_p if ar_p else np.nan)
            if not pd.isna(r_p) and r_p - r_c > 0.003:
                out.append(
                    f"**대손충당금 설정률**이 전기 {r_p:.2%}에서 당기 {r_c:.2%}로 하락함. "
                    f"채권 연령분석표를 입수하여 설정률 산정근거의 타당성을 확인할 것.")

    # 재고자산 대 매출원가
    inv_c, inv_p = F("재고", "당기금액"), F("재고", "전기금액")
    if inv_p and not pd.isna(cogs_g):
        inv_g = g(inv_c, inv_p)
        if inv_g - cogs_g > 0.10:
            out.append(
                f"**재고자산**이 {inv_g:+.1%} 변동하여 매출원가 증가율 {cogs_g:+.1%}를 상회함. "
                f"장기체화·진부화 재고 여부와 순실현가능가치 평가를 확인할 것.")

    # 매출총이익률
    if rev_c and rev_p:
        gp_c, gp_p = (rev_c - cogs_c) / rev_c, (rev_p - cogs_p) / rev_p
        if abs(gp_c - gp_p) >= 0.01:
            out.append(
                f"**매출총이익률**이 전기 {gp_p:.1%}에서 당기 {gp_c:.1%}로 "
                f"{(gp_c-gp_p)*100:+.1f}%p 변동함. 제품별·거래처별 원가율 분해가 필요함.")

    # 판관비 대 매출
    if sga_p and not pd.isna(rev_g):
        sga_g = g(sga_c, sga_p)
        if sga_g - rev_g > 0.05:
            out.append(
                f"**판매비와관리비**가 {sga_g:+.1%} 증가하여 매출 증가율을 상회함. "
                f"증가 기여 계정을 특정하고 일회성 여부를 확인할 것.")

    # 차입금·이자비용 정합성
    dbt_c = F("차입금", "당기금액") + F("사채", "당기금액")
    dbt_p = F("차입금", "전기금액") + F("사채", "전기금액")
    int_c, int_p = F("이자비용", "당기금액"), F("이자비용", "전기금액")
    if dbt_p and int_p:
        r_c, r_p = int_c / dbt_c if dbt_c else np.nan, int_p / dbt_p
        if not pd.isna(r_c) and abs(r_c - r_p) > 0.005:
            out.append(
                f"**차입금 평균이자율**이 전기 {r_p:.2%}에서 당기 {r_c:.2%}로 변동함. "
                f"차입약정서와 대조하여 이자비용의 완전성을 확인할 것.")

    # 나머지 유의계정
    named = "매출채권|대손충당|재고|매출원가|차입금|사채|이자비용"
    rest = flagged[~flagged["계정과목"].astype(str).str.contains(named, na=False)
                   & ~flagged["세부구분"].isin(["매출", "매출원가"])]
    for _, r in rest.head(4).iterrows():
        pct = "—" if pd.isna(r["증감률"]) else f"{r['증감률']:+.1%}"
        name = r["계정과목"]
        out.append(
            f"**{name}**{accounts.josa(name)} {r['증감액']:,.0f}원({pct}) 변동함. "
            f"증감 사유에 대한 경영진 문의 및 증빙 대사가 필요함.")

    return out or ["설정하신 임계치를 넘는 유의적 증감이 없습니다. 임계치를 낮춰 다시 확인해 보세요."]


def _export(df, flagged, ratios, comments, amt_th, pct_th):
    wb = Workbook()
    ws = wb.active
    ws.title = "증감분석"
    ws["A1"] = "재무제표 증감분석"
    ws["A1"].font = Font(FONT, size=14, bold=True)
    ws["A2"] = f"유의성 기준: 금액 {amt_th:,.0f}원 · 증감률 {pct_th}%"
    ws["A2"].font = Font(FONT, size=9, italic=True, color="6B7280")

    cols = ["세부구분", "계정과목", "전기금액", "당기금액", "증감액", "증감률"]
    for c, h in enumerate(cols, 1):
        cell = ws.cell(4, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment("center", "center")
        cell.border = BOX

    d = df.sort_values(["세부구분", "계정과목"])
    r = 5
    for _, row in d.iterrows():
        ws.cell(r, 1, str(row["세부구분"]))
        ws.cell(r, 2, str(row["계정과목"]))
        ws.cell(r, 3, float(row["전기금액"]))
        ws.cell(r, 4, float(row["당기금액"]))
        ws.cell(r, 5, f"=D{r}-C{r}")
        ws.cell(r, 6, f'=IF(C{r}=0,"",E{r}/ABS(C{r}))')
        for c in range(1, 7):
            ws.cell(r, c).border = BOX
            if row["유의"]:
                ws.cell(r, c).fill = FLAG
        for c in (3, 4, 5):
            ws.cell(r, c).number_format = "#,##0;(#,##0);-"
        ws.cell(r, 6).number_format = "0.0%;(0.0%);-"
        r += 1

    ws.cell(r, 2, "합계").font = Font(FONT, bold=True)
    for c in (3, 4):
        ws.cell(r, c, f"=SUM({get_column_letter(c)}5:{get_column_letter(c)}{r-1})")
        ws.cell(r, c).font = Font(FONT, bold=True)
        ws.cell(r, c).number_format = "#,##0"
    for c in range(1, 7):
        ws.cell(r, c).fill = SUB
        ws.cell(r, c).border = BOX

    ws.cell(r + 2, 1, "음영 표시 = 유의성 기준 초과 계정").font = Font(FONT, size=9, color="6B7280")
    for col, w in zip("ABCDEF", [16, 24, 18, 18, 18, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("비율분석")
    for c, h in enumerate(ratios.columns, 1):
        cell = ws2.cell(1, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
    for i, (_, row) in enumerate(ratios.iterrows(), 2):
        for c, v in enumerate(row, 1):
            cell = ws2.cell(i, c, v)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
    for col, w in zip("ABCDE", [18, 14, 14, 14, 30]):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("코멘트")
    ws3["A1"] = "분석 코멘트 초안"
    ws3["A1"].font = Font(FONT, size=12, bold=True)
    ws3["A2"] = "규칙 기반 자동 생성. 근거 확인 후 수정하여 사용할 것."
    ws3["A2"].font = Font(FONT, size=9, italic=True, color="6B7280")
    for i, c in enumerate(comments, 4):
        cell = ws3.cell(i, 1, "· " + c.replace("**", ""))
        cell.font = Font(FONT, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws3.row_dimensions[i].height = 32
    ws3.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
