"""LG-100 계정별원장 편집 — 계정마다 나뉜 시트를 하나의 통합 원장으로."""
import io
from datetime import date

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import ledger as L
from core.env import in_browser, size_guard
from core.theme import rule

FONT = "Arial"
HDR = PatternFill("solid", fgColor="16202E")
SUB = PatternFill("solid", fgColor="E8E6DF")
OK_F = PatternFill("solid", fgColor="D1FAE5")
NG_F = PatternFill("solid", fgColor="FEE2E2")
_t = Side(style="thin", color="C9C6BA")
_m = Side(style="medium", color="16202E")
BOX = Border(left=_t, right=_t, top=_t, bottom=_t)

# 편집본 서식 — 계정과목이 맨 앞, 차-대 열 포함
OUT_COLS = ["계정과목", "날짜", "적요", "거래처코드", "거래처명", "사업자등록번호",
            "차변금액", "대변금액", "차-대", "잔액", "계정코드", "계정명"]


def render():
    up = st.file_uploader("계정별원장 — 계정마다 시트가 나뉜 파일 (.xls / .xlsx)",
                          type=["xls", "xlsx"])
    if up is None:
        st.info("회계프로그램에서 뽑은 계정별원장을 그대로 올리시면 됩니다.\n\n"
                "계정과목이 시트 이름에만 있어도 열로 만들어 하나의 표로 합칩니다. "
                "[전기이월]·[월계]·[누계] 소계행은 걷어내되, "
                "적요에 쓰인 대괄호([잡코리아]유료서비스 등)는 건드리지 않습니다.")
        return

    lvl = size_guard(up)
    if lvl == "hard":
        st.error(
            f"파일이 {up.size/1024/1024:.0f}MB입니다. 브라우저에서 처리하기에는 너무 큽니다.\n\n"
            "이 크기는 PC에 설치해서 쓰시는 편이 훨씬 빠릅니다. "
            "홈 화면의 내려받기 안내를 참고해 주세요.")
        return
    if lvl == "warn":
        st.warning(
            f"파일이 {up.size/1024/1024:.0f}MB입니다. 브라우저에서는 몇 분 걸릴 수 있고, "
            "처리 중에는 화면이 멈춘 것처럼 보입니다. 탭을 닫지 말고 기다려 주세요.")

    with st.spinner("시트를 읽는 중입니다…"):
        try:
            sheets = L.read_xls_sheets(up)
        except Exception as e:
            st.error(f"파일을 읽지 못했습니다 — {e}")
            return

    if not sheets:
        st.error("계정별원장 형태의 시트를 찾지 못했습니다. "
                 "'날짜'와 '차변'이 같은 행에 있는 머리글이 필요합니다.")
        return

    st.success(f"시트 {len(sheets):,}개를 읽었습니다.")
    with st.expander("시트 인식 결과", expanded=False):
        info = pd.DataFrame([{"원본시트": s["name"], "계정코드": s["code"],
                              "계정명": s["account"], "행수": len(s["df"])}
                             for s in sheets])
        st.dataframe(info, use_container_width=True, height=280, hide_index=True)
        miss = info[info["계정코드"] == ""]
        if len(miss):
            st.warning(f"계정코드를 못 찾은 시트 {len(miss)}개 — 시트명이 그대로 계정과목이 됩니다.")

    rule()

    # ── 정리 기준 ────────────────────────────────────────────
    st.markdown('<div class="section-label">정리 기준</div>', unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    with c1:
        drop_sub = st.checkbox("소계행 제거", value=True,
                               help="[전기이월] [월계] [누계] 행을 뺍니다")
    with c2:
        drop_empty = st.checkbox("금액 0인 행 제거", value=True)
    with c3:
        sort_by = st.selectbox("정렬", ["계정코드 → 날짜", "날짜 → 계정코드", "원본 순서"])

    with st.spinner("시트를 합치는 중입니다…"):
        df, summary = L.combine(sheets, drop_subtotal=drop_sub, drop_empty=drop_empty)

    if df.empty:
        st.error("정리 후 남은 데이터가 없습니다.")
        return

    if sort_by == "계정코드 → 날짜":
        df = df.sort_values(["계정코드", "날짜"], kind="stable").reset_index(drop=True)
    elif sort_by == "날짜 → 계정코드":
        df = df.sort_values(["날짜", "계정코드"], kind="stable").reset_index(drop=True)

    raw_rows = sum(len(s["df"]) for s in sheets)
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("원본", f"{raw_rows:,}행")
    m2.metric("통합 후", f"{len(df):,}행", f"-{raw_rows-len(df):,}", delta_color="off")
    m3.metric("계정 수", f"{df['계정코드'].nunique():,}개")
    m4.metric("거래처 수", f"{df['거래처명'].replace('', pd.NA).nunique():,}개")

    # ── 산출물 정보 ──────────────────────────────────────────
    st.markdown('<div class="section-label">산출물 정보</div>', unsafe_allow_html=True)
    d1, d2, d3, d4 = st.columns(4)
    with d1:
        company = st.text_input("회사명", "마플")
    with d2:
        default_end = date(2025, 12, 31)
        try:
            dd = pd.to_datetime(df["날짜"], errors="coerce", format="mixed")
            if dd.notna().any():
                default_end = dd.max().date()
        except Exception:
            pass
        period_end = st.date_input("결산일", value=default_end)
    with d3:
        version = st.number_input("버전", min_value=1, value=1, step=1)
    with d4:
        prepared = st.text_input("작성자", "")

    fname = L.suggest_filename(company, pd.Timestamp(period_end), int(version))
    st.caption(f"파일명 → **{fname}**")

    rule()

    # ── 검증 ────────────────────────────────────────────────
    st.markdown('<div class="section-label">검증</div>', unsafe_allow_html=True)
    res = L.verify(df, summary, pd.Timestamp(period_end))

    v1, v2, v3 = st.columns(3)
    v1.metric("차변 합계", f"{res['dr']:,.0f}")
    v2.metric("대변 합계", f"{res['cr']:,.0f}")
    v3.metric("차이", f"{res['diff']:,.0f}")

    if abs(res["diff"]) < 1:
        st.success("차변과 대변이 일치합니다.")
    else:
        st.error(f"차변과 대변이 {res['diff']:,.0f}원 일치하지 않습니다. "
                 f"원장에서 빠진 계정이 있는지 확인해 주세요.")

    st.dataframe(res["checks"], use_container_width=True, hide_index=True)

    st.caption("원장 [누계]와 대조 — 원장에 찍힌 누계 금액과 편집 결과가 맞는지 계정별로 "
               "대조합니다. 소계행을 제대로 걷어냈는지 확인하는 가장 확실한 방법입니다.")

    detail = {
        "[누계] 불일치 계정": res["bad_tie"],
        "차·대 동시 기재": res["both"],
        "음수 금액": res["neg"],
        "날짜 인식 실패": res["bad_date"],
        "회계기간 범위 밖": res["out_range"],
    }
    shown = {k: v for k, v in detail.items() if len(v)}
    if shown:
        st.markdown('<div class="section-label">확인 대상</div>', unsafe_allow_html=True)
        for name, t in shown.items():
            with st.expander(f"{name} — {len(t):,}건"):
                st.dataframe(t.head(500), use_container_width=True,
                             height=min(300, 40 + 35 * len(t)))

    # ── 계정별 요약 ──────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">계정별 집계</div>', unsafe_allow_html=True)
    sm_view = summary[["계정코드", "계정명", "행수", "차변합계", "대변합계"]].copy()
    sm_view["차-대"] = sm_view["차변합계"] - sm_view["대변합계"]
    st.dataframe(
        sm_view.style.format({"차변합계": "{:,.0f}", "대변합계": "{:,.0f}", "차-대": "{:,.0f}"}),
        use_container_width=True, height=320, hide_index=True)

    # ── 결과 ────────────────────────────────────────────────
    rule()
    st.markdown('<div class="section-label">통합 원장</div>', unsafe_allow_html=True)
    cols = [c for c in OUT_COLS if c in df.columns]
    st.dataframe(
        df[cols].head(300).style.format(
            {"차변금액": "{:,.0f}", "대변금액": "{:,.0f}",
             "차-대": "{:,.0f}", "잔액": "{:,.0f}"}),
        use_container_width=True, height=380)
    if len(df) > 300:
        st.caption(f"화면에는 앞 300행만 표시했습니다. "
                   f"내려받는 파일에는 {len(df):,}행 전부 들어갑니다.")

    if st.button("엑셀 만들기", type="primary"):
        with st.spinner(f"{len(df):,}행을 쓰는 중입니다…"):
            xlsx = _export(df[cols], summary, res, {
                "company": company, "period_end": pd.Timestamp(period_end),
                "version": int(version), "prepared": prepared,
                "source": up.name, "raw_rows": raw_rows,
                "sheet_count": len(sheets),
            })
        st.download_button(f"{fname} 내려받기", xlsx, file_name=fname,
                           mime="application/vnd.openxmlformats-officedocument."
                                "spreadsheetml.sheet")


# ─────────────────────────────────────────────────────────────
def _export(df: pd.DataFrame, summary: pd.DataFrame, res: dict, meta: dict) -> bytes:
    wb = Workbook()

    # ── 검증 ────────────────────────────────────────────────
    ws = wb.active
    ws.title = "검증"
    ws["A1"] = f"{meta['company']} 계정별원장 통합 검증"
    ws["A1"].font = Font(FONT, size=14, bold=True)
    ws["A2"] = (f"결산일 {meta['period_end']:%Y-%m-%d} · v{meta['version']} · "
                f"작성 {meta['prepared'] or '―'} {date.today():%Y-%m-%d}")
    ws["A2"].font = Font(FONT, size=9, italic=True, color="6B7280")
    ws["A3"] = (f"원본 {meta['source']} · 시트 {meta['sheet_count']:,}개 · "
                f"{meta['raw_rows']:,}행 → 통합 {len(df):,}행")
    ws["A3"].font = Font(FONT, size=9, italic=True, color="6B7280")

    last = len(df) + 2  # 통합원장 시트의 마지막 데이터 행 (1행 합계, 2행 머리글)
    r = 5
    for c, h in enumerate(["구분", "금액"], 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")
    for label, f in [("차변 합계", f"=SUM('통합원장'!G3:G{last})"),
                     ("대변 합계", f"=SUM('통합원장'!H3:H{last})")]:
        r += 1
        ws.cell(r, 1, label).font = Font(FONT, size=10)
        ws.cell(r, 2, f).number_format = "#,##0;(#,##0);-"
        for c in (1, 2):
            ws.cell(r, c).border = BOX
    r += 1
    ws.cell(r, 1, "차이 (차변 - 대변)").font = Font(FONT, bold=True)
    ws.cell(r, 2, f"=B{r-2}-B{r-1}").font = Font(FONT, bold=True)
    ws.cell(r, 2).number_format = "#,##0;(#,##0);-"
    good = abs(res["diff"]) < 1
    for c in (1, 2):
        ws.cell(r, c).border = Border(top=_m, bottom=_m, left=_t, right=_t)
        ws.cell(r, c).fill = OK_F if good else NG_F
    r += 1
    ws.cell(r, 1, "판정").font = Font(FONT, bold=True)
    ws.cell(r, 2, f'=IF(ABS(B{r-1})<1,"일치","불일치")').font = Font(FONT, bold=True)
    for c in (1, 2):
        ws.cell(r, c).border = BOX
        ws.cell(r, c).fill = OK_F if good else NG_F

    r += 2
    ws.cell(r, 1, "검증 항목").font = Font(FONT, bold=True, size=11)
    r += 1
    for c, h in enumerate(["항목", "결과", "내용"], 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")
    for _, row in res["checks"].iterrows():
        r += 1
        g = row["결과"] in ("일치", "없음", "정상", "참고")
        for c, v in enumerate(row.tolist(), 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if c == 2:
                cell.fill = OK_F if g else NG_F
                cell.alignment = Alignment("center", "center")
    for col, w in zip("ABC", [24, 16, 60]):
        ws.column_dimensions[col].width = w

    # ── 통합원장 — 편집본 서식 (1행 합계, 2행 머리글) ──────────
    ws2 = wb.create_sheet("통합원장")
    cols = list(df.columns)
    i_dr = cols.index("차변금액") + 1
    i_cr = cols.index("대변금액") + 1
    Ldr, Lcr = get_column_letter(i_dr), get_column_letter(i_cr)

    # 1행: 전체 범위 합계 (SUBTOTAL 아님 — 필터에 영향받지 않는다)
    ws2.cell(1, i_dr, f"=SUM({Ldr}3:{Ldr}{last})")
    ws2.cell(1, i_cr, f"=SUM({Lcr}3:{Lcr}{last})")
    for c in (i_dr, i_cr):
        cell = ws2.cell(1, c)
        cell.font = Font(FONT, bold=True, size=11)
        cell.number_format = "#,##0"
        cell.fill = SUB
        cell.border = Border(top=_m, bottom=_m, left=_t, right=_t)

    for c, h in enumerate(cols, 1):
        cell = ws2.cell(2, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")

    num_idx = {cols.index(c) + 1 for c in ("차변금액", "대변금액", "차-대", "잔액")
               if c in cols}
    body = Font(FONT, size=10)
    for i, rec in enumerate(df.itertuples(index=False), 3):
        for c, v in enumerate(rec, 1):
            cell = ws2.cell(i, c, v)
            cell.font = body
            if c in num_idx:
                cell.number_format = "#,##0;(#,##0);-"

    ws2.freeze_panes = "A3"
    ws2.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{last}"
    widths = {"계정과목": 22, "날짜": 12, "적요": 34, "거래처코드": 11, "거래처명": 26,
              "사업자등록번호": 15, "차변금액": 16, "대변금액": 16, "차-대": 16,
              "잔액": 18, "계정코드": 10, "계정명": 18}
    for c, h in enumerate(cols, 1):
        ws2.column_dimensions[get_column_letter(c)].width = widths.get(h, 14)

    # ── 계정별 집계 ──────────────────────────────────────────
    ws3 = wb.create_sheet("계정별집계")
    sm = summary[["계정코드", "계정명", "행수", "차변합계", "대변합계",
                  "전기이월_차변", "전기이월_대변", "원장누계_차변", "원장누계_대변"]].copy()
    for c, h in enumerate(list(sm.columns) + ["차-대", "누계대조"], 1):
        cell = ws3.cell(1, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.border = BOX
        cell.alignment = Alignment("center", "center")
    for i, rec in enumerate(sm.itertuples(index=False), 2):
        for c, v in enumerate(rec, 1):
            cell = ws3.cell(i, c, None if pd.isna(v) else v)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if c >= 4:
                cell.number_format = "#,##0;(#,##0);-"
        ws3.cell(i, 10, f"=D{i}-E{i}").number_format = "#,##0;(#,##0);-"
        ws3.cell(i, 11, f'=IF(H{i}="","―",IF(ABS(D{i}+F{i}-H{i})<1,"일치","불일치"))')
        for c in (10, 11):
            ws3.cell(i, c).font = Font(FONT, size=10)
            ws3.cell(i, c).border = BOX

    t = len(sm) + 2
    ws3.cell(t, 2, "합계").font = Font(FONT, bold=True)
    for c in (4, 5):
        Lc = get_column_letter(c)
        ws3.cell(t, c, f"=SUM({Lc}2:{Lc}{t-1})")
        ws3.cell(t, c).font = Font(FONT, bold=True)
        ws3.cell(t, c).number_format = "#,##0"
        ws3.cell(t, c).fill = SUB
        ws3.cell(t, c).border = Border(top=_m, bottom=_m, left=_t, right=_t)
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:K{t-1}"
    for c, w in zip("ABCDEFGHIJK", [10, 20, 9, 17, 17, 15, 15, 17, 17, 17, 11]):
        ws3.column_dimensions[c].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
