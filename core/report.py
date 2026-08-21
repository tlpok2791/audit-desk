"""적출 결과를 감사조서 형식 엑셀로 내보낸다."""
import io
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR = PatternFill("solid", fgColor="16202E")
SUB = PatternFill("solid", fgColor="E8E6DF")
_thin = Side(style="thin", color="C9C6BA")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def build(results: dict, meta: dict) -> bytes:
    """
    results: {테스트명: {"df": DataFrame, "basis": 판단근거 문자열}}
    meta   : {"title", "period", "pop_count", "pop_amount", "notes": [str]}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"

    ws["A1"] = meta["title"]
    ws["A1"].font = Font(FONT, size=14, bold=True)
    ws["A2"] = (f"대상기간 {meta['period']} · 모집단 {meta['pop_count']:,}건 / "
                f"{meta['pop_amount']:,.0f}원 · 작성 {datetime.now():%Y-%m-%d %H:%M}")
    ws["A2"].font = Font(FONT, size=9, italic=True, color="6B7280")

    headers = ["테스트 항목", "적출 건수", "적출 금액", "건수 비율", "금액 비율", "판단 근거"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(4, c, h)
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment("center", "center")
        cell.border = BOX

    row = 5
    for name, res in results.items():
        d = res["df"]
        amt = float(d["차변금액"].sum()) if len(d) and "차변금액" in d else 0.0
        ws.cell(row, 1, name).font = Font(FONT, bold=True)
        ws.cell(row, 2, len(d))
        ws.cell(row, 3, amt)
        ws.cell(row, 4, f"=IF({meta['pop_count']}=0,0,B{row}/{meta['pop_count']})")
        ws.cell(row, 5, f"=IF({meta['pop_amount']}=0,0,C{row}/{meta['pop_amount']})")
        ws.cell(row, 6, res["basis"]).font = Font(FONT, size=9)
        for c in range(1, 7):
            ws.cell(row, c).border = BOX
        ws.cell(row, 2).number_format = "#,##0"
        ws.cell(row, 3).number_format = "#,##0;(#,##0);-"
        ws.cell(row, 4).number_format = "0.0%"
        ws.cell(row, 5).number_format = "0.0%"
        row += 1

    ws.cell(row, 1, "합계 (테스트 간 중복 포함)").font = Font(FONT, bold=True)
    ws.cell(row, 2, f"=SUM(B5:B{row-1})").font = Font(FONT, bold=True)
    ws.cell(row, 3, f"=SUM(C5:C{row-1})").font = Font(FONT, bold=True)
    ws.cell(row, 2).number_format = "#,##0"
    ws.cell(row, 3).number_format = "#,##0"
    for c in range(1, 7):
        ws.cell(row, c).fill = SUB
        ws.cell(row, c).border = BOX

    row += 2
    ws.cell(row, 1, "데이터 품질 메모").font = Font(FONT, bold=True, size=10)
    for note in meta.get("notes", []) or ["특이사항 없음"]:
        row += 1
        ws.cell(row, 1, f"· {note}").font = Font(FONT, size=9, color="6B7280")

    for col, width in zip("ABCDEF", [24, 12, 18, 12, 12, 60]):
        ws.column_dimensions[col].width = width

    for name, res in results.items():
        _detail_sheet(wb, name[:31], res["df"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _detail_sheet(wb: Workbook, title: str, df: pd.DataFrame):
    ws = wb.create_sheet(title)
    if df.empty:
        ws["A1"] = "적출 항목 없음"
        ws["A1"].font = Font(FONT, italic=True, color="6B7280")
        return

    out = df.copy()
    for c in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[c]):
            out[c] = out[c].dt.strftime("%Y-%m-%d %H:%M").str.replace(" 00:00", "", regex=False)

    for c, h in enumerate(out.columns, 1):
        cell = ws.cell(1, c, str(h))
        cell.font = Font(FONT, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment("center", "center")
        cell.border = BOX

    for i, (_, r) in enumerate(out.iterrows(), 2):
        for c, v in enumerate(r, 1):
            val = float(v) if isinstance(v, (int, float, np.number)) and not pd.isna(v) else (
                "" if pd.isna(v) else v)
            cell = ws.cell(i, c, val)
            cell.font = Font(FONT, size=10)
            cell.border = BOX
            if isinstance(val, float):
                cell.number_format = "#,##0"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(out.columns))}{len(out)+1}"
    for c in range(1, len(out.columns) + 1):
        longest = max([len(str(out.columns[c-1]))] +
                      [len(str(x)) for x in out.iloc[:, c-1].head(200)])
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 3, 11), 40)
